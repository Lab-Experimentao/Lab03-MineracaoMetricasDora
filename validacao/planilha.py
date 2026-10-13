"""Planilha de rotulagem da amostra-ouro, uma por avaliador (seção 6).

Só traz o que o avaliador precisa para julgar (links para o repositório, a release e
os commits). Nenhuma saída da heurística entra aqui, para não influenciar o rótulo.
"""

from __future__ import annotations

from pathlib import Path

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.worksheet import Worksheet

TIPOS_PROJETO = ("biblioteca/framework", "aplicação/serviço", "ferramenta CLI", "outro")
ENTREGAS_REAIS = ("sim", "não", "incerto")
CORRETIVA = ("sim", "não")

COLUNAS_REPOSITORIOS = ["repositorio", "url_repositorio", "tipo_projeto", "entregas_reais", "observacoes"]
COLUNAS_RELEASES = [
    "repositorio", "tag", "tag_anterior", "publicada_em", "url_release", "url_commits",
    "corretiva", "observacoes",
]

INSTRUCOES = [
    "Rotulagem da amostra-ouro (Lab03, seção 6 do enunciado)",
    "",
    "Preencha sozinho, sem consultar os demais avaliadores, e commite este arquivo (issue #17).",
    "Use só os valores das listas suspensas; deixe 'observacoes' para dúvidas e justificativas.",
    "",
    "Aba 'repositorios' (uma linha por repositório):",
    "  tipo_projeto: biblioteca/framework, aplicação/serviço, ferramenta CLI ou outro.",
    "  entregas_reais: as releases do repositório representam entregas reais ao usuário? sim, não ou incerto.",
    "",
    "Aba 'releases' (5 releases sorteadas por repositório):",
    "  corretiva: o objetivo principal da release é corrigir um defeito recente? sim ou não.",
    "  Leia as release notes (url_release) e a lista de commits desde a release anterior (url_commits).",
    "  url_commits fica vazio quando a release é a primeira da história do repositório.",
    "  A release anterior é a publicada logo antes, por data. Em projetos com várias linhas de versão",
    "  (ex.: v3.0.15 publicada depois de v4.1.1), ela pode ser de outra linha e a lista de commits não",
    "  reflete a release; nesse caso, julgue pelas release notes e anote em 'observacoes'.",
]


def _aba(planilha: Workbook, titulo: str, colunas: list[str], linhas: list[dict]) -> Worksheet:
    aba = planilha.create_sheet(titulo)
    aba.append(colunas)
    for celula in aba[1]:
        celula.font = Font(bold=True)
    aba.freeze_panes = "A2"
    for linha in linhas:
        aba.append([linha.get(c, "") for c in colunas])
    for indice, coluna in enumerate(colunas, start=1):
        if coluna.startswith("url"):
            for (celula,) in aba.iter_rows(min_row=2, min_col=indice, max_col=indice):
                if celula.value:
                    celula.hyperlink = celula.value
                    celula.style = "Hyperlink"
        largura = max([len(coluna)] + [len(str(l.get(coluna, ""))) for l in linhas])
        aba.column_dimensions[aba.cell(1, indice).column_letter].width = min(largura + 2, 60)
    return aba


def _lista(aba: Worksheet, colunas: list[str], coluna: str, valores: tuple[str, ...]) -> None:
    validacao = DataValidation(type="list", formula1='"' + ",".join(valores) + '"', allow_blank=True)
    validacao.error = "Escolha um valor da lista: " + ", ".join(valores)
    validacao.showErrorMessage = True
    letra = aba.cell(1, colunas.index(coluna) + 1).column_letter
    validacao.add(f"{letra}2:{letra}{max(aba.max_row, 2)}")
    aba.add_data_validation(validacao)


def gerar_planilha(destino: Path, amostra: pd.DataFrame) -> None:
    """Grava a planilha em branco a partir da amostra-ouro (saída de `sortear_releases`)."""
    planilha = Workbook()
    instrucoes = planilha.active
    instrucoes.title = "instrucoes"
    for texto in INSTRUCOES:
        instrucoes.append([texto])
    instrucoes.column_dimensions["A"].width = 110

    repositorios = [
        {"repositorio": repo, "url_repositorio": f"https://github.com/{repo}"}
        for repo in amostra["repositorio"].drop_duplicates()
    ]
    aba = _aba(planilha, "repositorios", COLUNAS_REPOSITORIOS, repositorios)
    _lista(aba, COLUNAS_REPOSITORIOS, "tipo_projeto", TIPOS_PROJETO)
    _lista(aba, COLUNAS_REPOSITORIOS, "entregas_reais", ENTREGAS_REAIS)

    aba = _aba(planilha, "releases", COLUNAS_RELEASES, amostra.to_dict("records"))
    _lista(aba, COLUNAS_RELEASES, "corretiva", CORRETIVA)

    destino.parent.mkdir(parents=True, exist_ok=True)
    planilha.save(destino)
