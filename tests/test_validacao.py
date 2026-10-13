from dataclasses import replace
from pathlib import Path

import pandas as pd
import pytest
from openpyxl import load_workbook

from pipeline.config import carregar_config
from validacao import amostra as am
from validacao import etapa
from validacao.planilha import COLUNAS_RELEASES, COLUNAS_REPOSITORIOS, gerar_planilha

CONFIG = carregar_config(Path(__file__).resolve().parent.parent / "config.yaml")


def compares_de(repos, n_releases=8):
    """Cadeia `release` com `n_releases` por repositório, mais linhas de outras cadeias."""
    linhas = []
    for repo in repos:
        for i in range(n_releases):
            anterior = f"v1.{i - 1}" if i else ""
            for cadeia in ("release", "release_prerelease", "tag"):
                linhas.append({
                    "repositorio": repo, "cadeia": cadeia, "tag": f"v1.{i}",
                    "publicada_em": f"2026-0{1 + i % 9}-10T12:00:00+00:00", "tag_anterior": anterior,
                })
        linhas.append({"repositorio": repo, "cadeia": "tag", "tag": "so-tag",
                       "publicada_em": "2026-01-01T00:00:00+00:00", "tag_anterior": "v1.0"})
    return pd.DataFrame(linhas)


def releases_de(compares):
    principal = compares[compares["cadeia"] == "release"]
    return pd.DataFrame({
        "repositorio": principal["repositorio"], "tag": principal["tag"],
        "publicada_em": principal["publicada_em"], "draft": "0", "prerelease": "0",
        "url": "https://github.com/" + principal["repositorio"] + "/releases/tag/" + principal["tag"],
        "notas": "Notas de " + principal["tag"],
    })


def commits_de(compares, por_release=2):
    principal = compares[(compares["cadeia"] == "release") & (compares["tag_anterior"] != "")]
    return pd.DataFrame([
        {"repositorio": r.repositorio, "tag_anterior": r.tag_anterior, "tag": r.tag, "titulo": f"{r.tag}: commit {j}"}
        for r in principal.itertuples() for j in range(por_release)
    ])


def metricas_de(repos):
    return pd.DataFrame({"repositorio": repos})


REPOS = [f"dono/r{i:02d}" for i in range(20)]


def test_sorteio_de_repositorios_e_reproduzivel_e_nao_depende_da_ordem():
    a = am.sortear_repositorios(metricas_de(REPOS), 6, 42)
    b = am.sortear_repositorios(metricas_de(list(reversed(REPOS))), 6, 42)
    assert a == b and len(set(a)) == 6 and a == sorted(a)
    assert am.sortear_repositorios(metricas_de(REPOS), 6, 7) != a


def test_amostra_maior_que_a_populacao_falha():
    with pytest.raises(ValueError, match="amostra final tem 3"):
        am.sortear_repositorios(metricas_de(REPOS[:3]), 6, 42)


def test_sorteia_k_releases_distintas_so_da_cadeia_principal():
    compares = compares_de(REPOS[:3])
    amostra = am.sortear_releases(compares, releases_de(compares), REPOS[:3], 5, 42)
    assert list(amostra.columns) == am.COLUNAS_AMOSTRA
    assert amostra.groupby("repositorio")["tag"].nunique().tolist() == [5, 5, 5]
    assert "so-tag" not in set(amostra["tag"])
    assert amostra["url_release"].str.startswith("https://github.com/dono/").all()


def test_sorteio_de_releases_e_reproduzivel():
    compares = compares_de(REPOS[:2])
    a = am.sortear_releases(compares, releases_de(compares), REPOS[:2], 5, 42)
    b = am.sortear_releases(compares.sample(frac=1, random_state=1), releases_de(compares), REPOS[:2], 5, 42)
    pd.testing.assert_frame_equal(a, b)


def test_repositorio_com_poucas_releases_falha():
    compares = compares_de(REPOS[:1], n_releases=3)
    with pytest.raises(ValueError, match="3 releases"):
        am.sortear_releases(compares, releases_de(compares), REPOS[:1], 5, 42)


def test_tag_repetida_fica_com_a_anterior_verdadeira():
    compares = pd.DataFrame([
        {"repositorio": "d/r", "cadeia": "release", "tag": "2954", "publicada_em": "2026-01-24T09:25:55+00:00",
         "tag_anterior": "2954"},
        {"repositorio": "d/r", "cadeia": "release", "tag": "2954", "publicada_em": "2026-01-24T09:25:55+00:00",
         "tag_anterior": "2953"},
    ])
    elegiveis = am.releases_elegiveis(compares)
    assert elegiveis[["tag", "tag_anterior"]].values.tolist() == [["2954", "2953"]]


def test_links_codificam_caracteres_especiais():
    assert am.url_commits("d/r", "v1.0", "pkg/v1.1+build") == (
        "https://github.com/d/r/compare/v1.0...pkg%2Fv1.1%2Bbuild"
    )
    assert am.url_commits("d/r", "", "v1.0") == ""
    assert am.url_release("d/r", "a b") == "https://github.com/d/r/releases/tag/a%20b"


def test_link_da_release_sem_url_na_coleta_e_montado():
    compares = compares_de(REPOS[:1])
    releases = releases_de(compares).iloc[0:0]
    amostra = am.sortear_releases(compares, releases, REPOS[:1], 5, 42)
    assert amostra["url_release"].str.contains("/releases/tag/v1.").all()


def test_resumos_cortam_texto_longo():
    assert am.resumir_notas("") == "(release sem notas)"
    assert am.resumir_notas("abc") == "abc"
    cortado = am.resumir_notas("x" * 1500)
    assert cortado.startswith("x" * 1000) and "+500 caracteres" in cortado
    assert am.resumir_titulos([]) == "(nenhum commit novo em relação à anterior)"
    titulos = am.resumir_titulos([f"t{i}" for i in range(25)])
    assert titulos.count("\n- ") == 19 and "+5 commits" in titulos


def test_textos_vem_das_notas_e_dos_commits_da_release():
    compares = compares_de(REPOS[:1])
    compares["status"] = "ok"
    compares.loc[compares["tag"] == "v1.0", "status"] = "primeira"
    compares.loc[(compares["cadeia"] == "release") & (compares["tag"] == "v1.3"), "status"] = "404"
    releases = releases_de(compares)
    amostra = pd.DataFrame([
        {"repositorio": REPOS[0], "tag": t, "tag_anterior": a, "publicada_em": releases.set_index("tag").loc[t, "publicada_em"]}
        for t, a in [("v1.0", ""), ("v1.2", "v1.1"), ("v1.3", "v1.2")]
    ])
    textos = am.com_textos(amostra, compares, releases, commits_de(compares))
    assert textos["notas_release"].tolist() == ["Notas de v1.0", "Notas de v1.2", "Notas de v1.3"]
    primeira, normal, com_erro = textos["titulos_commits"]
    assert "primeira release" in primeira
    assert normal == "- v1.2: commit 0\n- v1.2: commit 1"
    assert "compare retornou 404" in com_erro


def test_planilha_tem_so_as_colunas_de_rotulagem_e_listas(tmp_path):
    compares = compares_de(REPOS[:2])
    compares["status"] = "ok"
    releases = releases_de(compares)
    amostra = am.sortear_releases(compares, releases, REPOS[:2], 5, 42)
    amostra = am.com_textos(amostra, compares, releases, commits_de(compares))
    amostra.loc[0, "notas_release"] = "=SOMA(A1)\x07 sem controle"
    destino = tmp_path / "rotulagem_x.xlsx"
    gerar_planilha(destino, amostra)

    planilha = load_workbook(destino)
    assert planilha.sheetnames == ["instrucoes", "repositorios", "releases"]
    repos = planilha["repositorios"]
    releases = planilha["releases"]
    assert [c.value for c in repos[1]] == COLUNAS_REPOSITORIOS
    assert [c.value for c in releases[1]] == COLUNAS_RELEASES
    assert repos.max_row == 3 and releases.max_row == 11
    # Colunas de rótulo em branco, com lista suspensa.
    assert all(c.value is None for c in releases["E"][1:])
    listas = {str(v.sqref): v.formula1 for v in releases.data_validations.dataValidation}
    assert listas == {"E2:E11": '"sim,não"'}
    assert len(repos.data_validations.dataValidation) == 2
    assert releases["H2"].hyperlink.target == releases["H2"].value
    # Texto que parece fórmula fica como texto, e caracteres de controle são removidos.
    assert releases["F2"].data_type == "s" and releases["F2"].value == "=SOMA(A1) sem controle"


def test_etapa_grava_amostra_e_nao_sobrescreve_planilha(tmp_path):
    compares = compares_de(REPOS)
    brutos = tmp_path / "brutos"
    brutos.mkdir()
    compares["status"] = "ok"
    compares.to_csv(brutos / "compares.csv", index=False)
    releases_de(compares).to_csv(brutos / "releases.csv", index=False)
    commits_de(compares).to_csv(brutos / "commits.csv", index=False)
    metricas_de(REPOS).to_csv(tmp_path / "metricas.csv", index=False)
    config = replace(CONFIG, saida_dir=tmp_path, validacao_manual=replace(
        CONFIG.validacao_manual, n_repositorios=4, avaliadores=("ana", "bia")))

    preenchida = etapa.arquivo_rotulagem(tmp_path / "validacao", "bia")
    preenchida.parent.mkdir()
    preenchida.write_bytes(b"rotulos")
    etapa.executar(config)

    amostra = pd.read_csv(tmp_path / "validacao" / etapa.ARQUIVO_AMOSTRA)
    assert len(amostra) == 4 * 5
    assert load_workbook(etapa.arquivo_rotulagem(tmp_path / "validacao", "ana")).sheetnames[0] == "instrucoes"
    assert preenchida.read_bytes() == b"rotulos"
