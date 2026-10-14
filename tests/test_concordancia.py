from dataclasses import replace
from pathlib import Path

import pandas as pd
import pytest
from openpyxl import load_workbook

from pipeline.config import carregar_config
from validacao import concordancia as co
from validacao import etapa
from validacao.planilha import gerar_planilha

CONFIG = carregar_config(Path(__file__).resolve().parent.parent / "config.yaml")

REPOS = ["d/a", "d/b", "d/c", "d/d"]


def abas(tipos, entregas, corretivas):
    """Abas de uma planilha preenchida: um repositório por tipo e uma release por rótulo de corretiva."""
    repositorios = pd.DataFrame({
        "repositorio": REPOS[:len(tipos)], "url_repositorio": "", "tipo_projeto": tipos,
        "entregas_reais": entregas, "observacoes": "",
    })
    releases = pd.DataFrame({
        "repositorio": "d/a", "tag": [f"v{i}" for i in range(len(corretivas))], "tag_anterior": "",
        "publicada_em": "2026-01-01T00:00:00+00:00", "corretiva": corretivas, "observacoes": "",
    })
    return {"repositorios": repositorios, "releases": releases}


BIB, APP, CLI = "biblioteca/framework", "aplicação/serviço", "ferramenta CLI"


def tres_avaliadores():
    return {
        "ana": abas([BIB, APP, CLI], ["sim", "sim", "sim"], ["sim", "sim", "não", "sim"]),
        "bia": abas([BIB, APP, APP], ["sim", "não", "sim"], ["sim", "sim", "não", "não"]),
        "caio": abas([BIB, CLI, "outro"], ["sim", "incerto", "sim"], ["sim", "não", "não", "não"]),
    }


def sem_desempates():
    return pd.DataFrame(columns=co.COLUNAS_DESEMPATES)


def test_kappa_confere_com_a_conta_a_mao():
    # P̄ = (1 + 1/3 + 1 + 1/3) / 4 = 2/3; Pe = 0,5² + 0,5² = 0,5; kappa = (2/3 - 0,5) / 0,5 = 1/3.
    rotulos = pd.DataFrame([["a", "a", "a"], ["a", "a", "b"], ["b", "b", "b"], ["a", "b", "b"]])
    assert co.kappa(rotulos) == pytest.approx(1 / 3)
    assert co.concordancia_observada(rotulos) == pytest.approx(2 / 3)


def test_kappa_com_concordancia_perfeita_e_um():
    assert co.kappa(pd.DataFrame([["a", "a"], ["b", "b"]])) == pytest.approx(1.0)


@pytest.mark.parametrize("valor, esperado", [
    (-0.1, "pobre"), (0.0, "leve"), (0.2, "leve"), (0.35, "razoável"), (0.6, "moderada"),
    (0.61, "substancial"), (0.95, "quase perfeita"), (float("nan"), "indefinido"),
])
def test_interpretacao_de_landis_e_koch(valor, esperado):
    assert co.interpretar(valor) == esperado


def test_consenso_unanime_maioria_e_pendente():
    assert co.consenso(pd.Series(["sim", "sim", "sim"]), None) == ("sim", co.UNANIME)
    assert co.consenso(pd.Series(["sim", "não", "sim"]), None) == ("sim", co.MAIORIA)
    assert co.consenso(pd.Series(["a", "b", "c"]), None) == ("", co.PENDENTE)
    # Com dois avaliadores, 1 a 1 não é maioria.
    assert co.consenso(pd.Series(["a", "b"]), None) == ("", co.PENDENTE)


def test_decisao_de_reuniao_prevalece_ate_sobre_a_maioria():
    assert co.consenso(pd.Series(["a", "b", "c"]), "b") == ("b", co.REUNIAO)
    assert co.consenso(pd.Series(["sim", "sim", "não"]), "não") == ("não", co.REUNIAO)


def test_problemas_aponta_vazios_e_valores_fora_da_lista():
    planilha = abas([BIB, "lib"], ["sim", ""], ["sim"])
    assert co.problemas(planilha) == [
        "repositorios linha 3: tipo_projeto = 'lib' fora da lista",
        "repositorios linha 3: entregas_reais vazio",
    ]
    assert co.problemas(abas([BIB], ["sim"], ["não"])) == []


def test_calcular_resume_cada_dimensao_e_monta_o_consenso():
    resumo, consensos = co.calcular(tres_avaliadores(), sem_desempates())
    linhas = resumo.set_index("dimensao")
    assert list(resumo.columns) == co.COLUNAS_CONCORDANCIA
    assert linhas.loc["tipo_projeto", ["itens", "unanimes", "maioria", "sem_maioria", "pendentes"]].tolist() == [3, 1, 1, 1, 1]
    assert linhas.loc["corretiva", ["unanimes", "maioria", "sem_maioria"]].tolist() == [2, 2, 0]
    assert linhas.loc["corretiva", "avaliadores"] == 3

    repos = consensos["repositorios"]
    assert repos["tipo_projeto"].tolist() == [BIB, APP, ""]
    assert repos["tipo_projeto_regra"].tolist() == [co.UNANIME, co.MAIORIA, co.PENDENTE]
    assert repos["entregas_reais"].tolist() == ["sim", "", "sim"]
    assert repos["tipo_projeto_bia"].tolist() == [BIB, APP, APP]

    releases = consensos["releases"]
    assert list(releases.columns[:4]) == ["repositorio", "tag", "tag_anterior", "publicada_em"]
    assert releases["corretiva"].tolist() == ["sim", "sim", "não", "não"]


def test_desempates_resolvem_os_pendentes():
    desempates = pd.DataFrame([
        {"dimensao": "tipo_projeto", "repositorio": "d/c", "tag": "", "rotulo": CLI, "justificativa": "reunião"},
        {"dimensao": "corretiva", "repositorio": "d/a", "tag": "v1", "rotulo": "não", "justificativa": "reunião"},
    ])
    resumo, consensos = co.calcular(tres_avaliadores(), desempates)
    assert consensos["repositorios"].loc[2, ["tipo_projeto", "tipo_projeto_regra"]].tolist() == [CLI, co.REUNIAO]
    assert consensos["releases"].loc[1, ["corretiva", "corretiva_regra"]].tolist() == ["não", co.REUNIAO]
    linhas = resumo.set_index("dimensao")
    # As contagens de concordância continuam as da votação, sem as reuniões.
    assert linhas.loc["tipo_projeto", ["sem_maioria", "decididos_em_reuniao", "pendentes"]].tolist() == [1, 1, 0]
    assert linhas.loc["corretiva", ["maioria", "decididos_em_reuniao"]].tolist() == [2, 1]


@pytest.mark.parametrize("desempate, mensagem", [
    ({"dimensao": "tipo_projeto", "repositorio": "d/z", "tag": "", "rotulo": CLI}, "não está na amostra"),
    ({"dimensao": "tipo_projeto", "repositorio": "d/c", "tag": "", "rotulo": "lib"}, "fora de"),
    ({"dimensao": "corretiva", "repositorio": "d/a", "tag": "v9", "rotulo": "sim"}, "não está na amostra"),
    ({"dimensao": "tipo", "repositorio": "d/c", "tag": "", "rotulo": CLI}, "dimensão desconhecida"),
])
def test_desempate_invalido_falha(desempate, mensagem):
    with pytest.raises(co.ErroRotulos, match=mensagem):
        co.calcular(tres_avaliadores(), pd.DataFrame([{**desempate, "justificativa": ""}]))


def test_desempate_repetido_falha():
    linha = {"dimensao": "tipo_projeto", "repositorio": "d/c", "tag": "", "rotulo": CLI, "justificativa": ""}
    with pytest.raises(co.ErroRotulos, match="mais de uma vez"):
        co.calcular(tres_avaliadores(), pd.DataFrame([linha, linha]))


def test_planilhas_com_linhas_diferentes_falham():
    planilhas = tres_avaliadores()
    planilhas["bia"]["releases"] = planilhas["bia"]["releases"].iloc[::-1].reset_index(drop=True)
    with pytest.raises(co.ErroRotulos, match="mesmas linhas"):
        co.calcular(planilhas, sem_desempates())


def test_arquivo_de_desempates(tmp_path):
    assert co.ler_desempates(tmp_path / "nao_existe.csv").empty
    caminho = tmp_path / "desempates.csv"
    caminho.write_text("dimensao,repositorio,rotulo\n", encoding="utf-8")
    with pytest.raises(co.ErroRotulos, match="justificativa"):
        co.ler_desempates(caminho)


# --- etapa, com planilhas .xlsx de verdade ---

def amostra_ouro():
    return pd.DataFrame({
        "repositorio": ["d/a", "d/a", "d/b", "d/b"], "tag": ["v1", "v2", "v1", "v2"],
        "tag_anterior": ["", "v1", "", "v1"], "publicada_em": "2026-01-01T00:00:00+00:00",
        "url_release": "", "url_commits": "", "notas_release": "notas", "titulos_commits": "- commit",
    })


def preencher(caminho, tipos, entregas, corretivas):
    planilha = load_workbook(caminho)
    for linha, (tipo, entrega) in enumerate(zip(tipos, entregas), start=2):
        planilha["repositorios"].cell(linha, 3, tipo)
        planilha["repositorios"].cell(linha, 4, f" {entrega} ")  # espaço digitado por engano
    for linha, corretiva in enumerate(corretivas, start=2):
        planilha["releases"].cell(linha, 5, corretiva)
    planilha.save(caminho)


def test_etapa_grava_concordancia_e_consenso_quando_todas_estao_completas(tmp_path, caplog):
    config = replace(CONFIG, validacao_manual=replace(CONFIG.validacao_manual, avaliadores=("ana", "bia", "caio")))
    rotulos = {
        "ana": ([BIB, APP], ["sim", "sim"], ["sim", "não", "não", "sim"]),
        "bia": ([BIB, CLI], ["sim", "não"], ["sim", "não", "sim", "sim"]),
        "caio": ([BIB, APP], ["sim", "incerto"], ["sim", "não", "não", "não"]),
    }
    for avaliador, valores in rotulos.items():
        caminho = etapa.arquivo_rotulagem(tmp_path, avaliador)
        gerar_planilha(caminho, amostra_ouro())
        if avaliador != "caio":
            preencher(caminho, *valores)

    etapa.calcular_concordancia(config, tmp_path)
    assert "rotulagem de caio incompleta" in caplog.text
    assert not (tmp_path / etapa.ARQUIVO_CONCORDANCIA).exists()

    preencher(etapa.arquivo_rotulagem(tmp_path, "caio"), *rotulos["caio"])
    etapa.calcular_concordancia(config, tmp_path)
    resumo = pd.read_csv(tmp_path / etapa.ARQUIVO_CONCORDANCIA)
    assert resumo["dimensao"].tolist() == ["tipo_projeto", "entregas_reais", "corretiva"]
    consenso = pd.read_csv(tmp_path / "consenso_repositorios.csv", keep_default_na=False)
    assert consenso["entregas_reais"].tolist() == ["sim", ""]
    assert "sem maioria em entregas_reais de d/b" in caplog.text
    assert pd.read_csv(tmp_path / etapa.ARQUIVO_DESEMPATES).columns.tolist() == co.COLUNAS_DESEMPATES

    (tmp_path / etapa.ARQUIVO_DESEMPATES).write_text(
        "dimensao,repositorio,tag,rotulo,justificativa\nentregas_reais,d/b,,sim,decidido em reunião\n",
        encoding="utf-8",
    )
    etapa.calcular_concordancia(config, tmp_path)
    consenso = pd.read_csv(tmp_path / "consenso_repositorios.csv", keep_default_na=False)
    assert consenso[["entregas_reais", "entregas_reais_regra"]].values.tolist() == [
        ["sim", co.UNANIME], ["sim", co.REUNIAO],
    ]
    assert pd.read_csv(tmp_path / "consenso_releases.csv")["corretiva"].tolist() == ["sim", "não", "não", "sim"]


def test_etapa_com_todas_as_planilhas_nao_precisa_dos_brutos(tmp_path):
    config = replace(CONFIG, saida_dir=tmp_path,
                     validacao_manual=replace(CONFIG.validacao_manual, avaliadores=("ana", "bia")))
    for avaliador in ("ana", "bia"):
        caminho = etapa.arquivo_rotulagem(tmp_path / "validacao", avaliador)
        gerar_planilha(caminho, amostra_ouro())
        preencher(caminho, [BIB, APP], ["sim", "sim"], ["sim", "não", "não", "sim"])
    etapa.executar(config)  # sem metricas.csv nem brutos/
    assert pd.read_csv(tmp_path / "validacao" / etapa.ARQUIVO_CONCORDANCIA)["unanimes"].tolist() == [2, 2, 4]
